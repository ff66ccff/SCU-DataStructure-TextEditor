"""Extract course-guide paragraphs and tables (requires python-docx)."""
import argparse
from pathlib import Path
import docx

root = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("source", nargs="?", type=Path,
                    default=root / "reference/数据结构文本编辑器实验指导书_学生版.docx")
parser.add_argument("--output", type=Path, default=root / "reference/guide-extract.md")
args = parser.parse_args()
document = docx.Document(args.source)
lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
lines.extend(["", f"=== TABLES: {len(document.tables)}"])
for index, table in enumerate(document.tables):
    lines.append(f"--- table {index} ---")
    for row in table.rows:
        lines.append(" | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))
args.output.write_text("\n".join(lines), encoding="utf-8")
print(args.output)
